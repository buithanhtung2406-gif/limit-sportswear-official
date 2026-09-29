from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


def number(value):
    try:
        return float(str(value).replace(",", "")) if value not in (None, "", "-") else 0
    except (TypeError, ValueError):
        return 0


def text(value):
    return str(value).strip() if value not in (None, "-") else ""


source = Path(sys.argv[1])
output = Path(sys.argv[2])
workbook = load_workbook(source, data_only=True, read_only=True)

shopee_rows = []
last_date = None
last_buyer = ""
for row in workbook["SHOPEE"].iter_rows(values_only=True):
    if isinstance(row[0], datetime):
        last_date = row[0]
        last_buyer = text(row[1])
    elif row[0] not in (None, ""):
        continue
    if not last_date or not text(row[5]):
        continue
    if last_date.strftime("%Y-%m") not in {"2026-07", "2026-08"}:
        continue
    shopee_rows.append({
        "orderDate": last_date.isoformat(timespec="seconds"), "buyer": text(row[1]) or last_buyer,
        "sku": text(row[2]), "color": text(row[3]), "size": text(row[4]), "productName": text(row[5]),
        "orderStatus": text(row[6]), "trackingCode": text(row[7]), "originalPrice": number(row[8]),
        "sellerDiscount": number(row[9]), "shopeeDiscount": number(row[10]), "totalSellerDiscount": number(row[11]),
        "promotionalPrice": number(row[12]), "quantity": number(row[13]), "buyerProductPayment": number(row[14]),
        "orderValue": number(row[15]), "shopVoucher": number(row[16]), "shopeeVoucher": number(row[17]),
        "shopeeCoins": number(row[18]), "expectedShippingFee": number(row[19]), "buyerShippingFee": number(row[20]),
        "shopeeShippingSubsidy": number(row[21]), "returnShippingFee": number(row[22]), "buyerTotalPaid": number(row[23]),
        "paymentMethod": text(row[24]), "fixedFee": number(row[25]), "serviceFee": number(row[26]),
        "transactionFee": number(row[27]), "province": text(row[28]),
    })

spx_rows = []
for row in workbook["SPX"].iter_rows(values_only=True):
    if isinstance(row[0], datetime) or not text(row[0]) or row[0] == "Mã vận đơn" or text(row[1]).upper() in {"TỔNG", "TÔNG"}:
        continue
    spx_rows.append({
        "trackingCode": text(row[0]), "sku": text(row[1]), "color": text(row[2]), "size": text(row[3]),
        "province": text(row[4]), "recipient": text(row[5]), "codAmount": number(row[6]),
        "failedDelivery": number(row[7]), "shippingFee": number(row[8]), "total": number(row[9]),
    })

inventory_rows = []
for sheet_name, snapshot_date in [("Tồn kho", "2026-09-03"), ("Bản sao của Trang tính9", "2026-09-05")]:
    for row in workbook[sheet_name].iter_rows(min_row=4, max_row=18, values_only=True):
        if not any(value not in (None, "") for value in row[:8]):
            continue
        size_s, size_m, size_l = number(row[3]), number(row[4]), number(row[5])
        inventory_rows.append({
            "snapshotDate": snapshot_date, "sku": text(row[0]), "itemType": text(row[1]), "color": text(row[2]),
            "sizeS": size_s, "sizeM": size_m, "sizeL": size_l,
            "total": number(row[6]) or size_s + size_m + size_l, "note": text(row[7]),
        })

catalog_rows = []
for row in workbook["mã sản phẩm"].iter_rows(min_row=2, values_only=True):
    if not text(row[0]):
        continue
    catalog_rows.append({"sku": text(row[0]), "name": text(row[1]), "price": number(row[2]), "costPrice": 0})

payload = {"action": "importWorkbookData", "shopeeRows": shopee_rows, "spxRows": spx_rows, "inventoryRows": inventory_rows, "catalogRows": catalog_rows}
output.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
print(json.dumps({"shopee": len(shopee_rows), "spx": len(spx_rows), "inventory": len(inventory_rows), "catalog": len(catalog_rows)}, ensure_ascii=False))
